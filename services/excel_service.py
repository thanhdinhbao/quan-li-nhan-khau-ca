import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PyQt5.QtCore import QThread, pyqtSignal
from datetime import datetime
from typing import List, Dict, Tuple, Any

from database.db_manager import DatabaseManager
from database.models import Person, Area
from controllers.person_controller import PersonController
from controllers.area_controller import AreaController
from utils.formatters import parse_date_to_iso, vietnamese_unaccent, format_date_vn, mask_id_card, mask_phone
from utils.validators import validate_full_name, validate_id_card, validate_dob, validate_phone
from utils.logger import log_info, log_error

class ExcelImportWorker(QThread):
    progress = pyqtSignal(int, int) # (current, total)
    finished_import = pyqtSignal(dict) # Summary dictionary
    error_occurred = pyqtSignal(str)

    def __init__(self, file_path: str, sheet_name: str, column_mapping: dict, update_existing: bool = False, user_id: int = None):
        super().__init__()
        self.file_path = file_path
        self.sheet_name = sheet_name
        self.column_mapping = column_mapping
        self.update_existing = update_existing
        self.user_id = user_id

    def run(self):
        db = DatabaseManager()
        area_ctrl = AreaController()
        person_ctrl = PersonController()

        areas = area_ctrl.get_all_areas()
        area_map = {vietnamese_unaccent(a.name): a.id for a in areas}
        default_area_id = areas[0].id if areas else 1

        try:
            # Read Excel file treating CCCD and Phone as string to preserve leading zeros
            df = pd.read_excel(self.file_path, sheet_name=self.sheet_name, dtype=str)
            df = df.fillna("")

            total_rows = len(df)
            success_cnt = 0
            skipped_cnt = 0
            error_cnt = 0
            error_details = []

            conn = db.get_connection()
            cursor = conn.cursor()
            conn.execute("BEGIN TRANSACTION;")

            start_time = datetime.now()

            cursor.execute("SELECT MAX(id) as max_id FROM persons")
            start_row = cursor.fetchone()
            max_id = start_row["max_id"] if start_row and start_row["max_id"] else 0

            cursor.execute("SELECT MAX(CAST(SUBSTR(person_code, 3) AS INTEGER)) as max_code FROM persons WHERE person_code LIKE 'NK%'")
            start_code_row = cursor.fetchone()
            max_code = start_code_row["max_code"] if start_code_row and start_code_row["max_code"] else 0

            current_max_num = max(max_id, max_code)

            for idx, row in df.iterrows():
                row_num = idx + 2 # Excel 1-based + 1 header
                self.progress.emit(idx + 1, total_rows)

                # Extract mapped fields
                name = str(row.get(self.column_mapping.get("full_name", ""), "")).strip()
                cccd = str(row.get(self.column_mapping.get("id_card_number", ""), "")).strip()
                dob_raw = str(row.get(self.column_mapping.get("dob", ""), "")).strip()
                gender = str(row.get(self.column_mapping.get("gender", ""), "")).strip() or "Nam"
                phone = str(row.get(self.column_mapping.get("phone", ""), "")).strip()
                hometown = str(row.get(self.column_mapping.get("hometown", ""), "")).strip()
                area_raw = str(row.get(self.column_mapping.get("area_name", ""), "")).strip()
                addr = str(row.get(self.column_mapping.get("permanent_address", ""), "")).strip()

                if not name:
                    error_cnt += 1
                    error_details.append(f"Dòng {row_num}: Họ và tên bị trống.")
                    continue

                valid, msg = validate_full_name(name)
                if not valid:
                    error_cnt += 1
                    error_details.append(f"Dòng {row_num}: {msg}")
                    continue

                dob_iso = parse_date_to_iso(dob_raw)

                # Area mapping
                area_id = default_area_id
                if area_raw:
                    norm_area = vietnamese_unaccent(area_raw)
                    if norm_area in area_map:
                        area_id = area_map[norm_area]

                # Check duplicate CCCD
                duplicate = False
                if cccd:
                    cursor.execute("SELECT id FROM persons WHERE id_card_number = ? AND is_deleted = 0", (cccd,))
                    existing_row = cursor.fetchone()
                    if existing_row:
                        duplicate = True
                        if not self.update_existing:
                            skipped_cnt += 1
                            continue

                while True:
                    current_max_num += 1
                    candidate_code = f"NK{current_max_num:06d}"
                    cursor.execute("SELECT 1 FROM persons WHERE person_code = ?", (candidate_code,))
                    if not cursor.fetchone():
                        code = candidate_code
                        break

                norm_name = vietnamese_unaccent(name)

                if duplicate and self.update_existing:
                    cursor.execute("""
                        UPDATE persons SET
                            full_name = ?, normalized_name = ?, dob = ?, gender = ?, phone = ?, hometown = ?, permanent_address = ?, area_id = ?, updated_at = CURRENT_TIMESTAMP, updated_by = ?
                        WHERE id_card_number = ?
                    """, (name, norm_name, dob_iso, gender, phone, hometown, addr, area_id, self.user_id, cccd))
                    success_cnt += 1
                else:
                    cursor.execute("""
                        INSERT INTO persons (
                            person_code, full_name, normalized_name, dob, gender, id_card_number, phone, hometown, permanent_address, area_id, created_by, updated_by
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (code, name, norm_name, dob_iso, gender, cccd, phone, hometown, addr, area_id, self.user_id, self.user_id))
                    success_cnt += 1

            conn.commit()
            conn.close()

            elapsed = (datetime.now() - start_time).total_seconds()

            summary = {
                "total": total_rows,
                "success": success_cnt,
                "skipped": skipped_cnt,
                "error": error_cnt,
                "elapsed_seconds": round(elapsed, 2),
                "error_details": error_details
            }
            self.finished_import.emit(summary)
        except Exception as e:
            log_error(f"Excel import error: {e}", exc_info=True)
            self.error_occurred.emit(str(e))

class ExcelExportService:
    @staticmethod
    def export_persons_to_excel(file_path: str, persons: List[Person], has_full_permission: bool = True) -> bool:
        """Exports persons list to stylized OpenPyXL Excel spreadsheet."""
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Danh sách Nhân khẩu"
            ws.views.sheetView[0].showGridLines = True

            # Colors & Styles (GOV Red Header)
            header_fill = PatternFill(start_color="A6192E", end_color="A6192E", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            title_font = Font(name="Segoe UI", size=14, bold=True, color="A6192E")
            sub_font = Font(name="Segoe UI", size=10, italic=True, color="555555")
            data_font = Font(name="Segoe UI", size=10)
            center_align = Alignment(horizontal="center", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")

            thin_border = Border(
                left=Side(style='thin', color='D3D3D3'),
                right=Side(style='thin', color='D3D3D3'),
                top=Side(style='thin', color='D3D3D3'),
                bottom=Side(style='thin', color='D3D3D3')
            )

            # Title Rows
            ws.merge_cells("A1:J1")
            ws["A1"] = "HỆ THỐNG QUẢN LÝ NHÂN KHẨU CẤP XÃ - CÔNG AN XÃ PHÚ LƯƠNG"
            ws["A1"].font = title_font
            ws["A1"].alignment = left_align

            ws.merge_cells("A2:J2")
            ws["A2"] = f"DANH SÁCH BÁO CÁO NHÂN KHẨU (Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')})"
            ws["A2"].font = sub_font
            ws["A2"].alignment = left_align

            # Headers
            headers = [
                "STT", "Mã nhân khẩu", "Họ và tên", "Ngày sinh", "Giới tính",
                "Số CCCD/CMND", "Số điện thoại", "Khu vực", "Địa chỉ thường trú", "Trạng thái cư trú"
            ]

            header_row = 4
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=header_row, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            # Write Data
            for r_idx, p in enumerate(persons, 1):
                row_num = header_row + r_idx
                ws.cell(row=row_num, column=1, value=r_idx).alignment = center_align
                ws.cell(row=row_num, column=2, value=p.person_code).alignment = center_align
                ws.cell(row=row_num, column=3, value=p.full_name).alignment = left_align
                ws.cell(row=row_num, column=4, value=format_date_vn(p.dob)).alignment = center_align
                ws.cell(row=row_num, column=5, value=p.gender).alignment = center_align
                
                # Format CCCD and Phone as explicit TEXT string
                cccd_val = mask_id_card(p.id_card_number, has_full_permission)
                cell_cccd = ws.cell(row=row_num, column=6, value=cccd_val)
                cell_cccd.number_format = '@'
                cell_cccd.alignment = center_align

                phone_val = mask_phone(p.phone, has_full_permission)
                cell_phone = ws.cell(row=row_num, column=7, value=phone_val)
                cell_phone.number_format = '@'
                cell_phone.alignment = center_align

                ws.cell(row=row_num, column=8, value=p.area_name or "").alignment = left_align
                ws.cell(row=row_num, column=9, value=p.permanent_address or "").alignment = left_align
                ws.cell(row=row_num, column=10, value=p.residence_status or "").alignment = center_align

                for col_idx in range(1, 11):
                    c = ws.cell(row=row_num, column=col_idx)
                    c.font = data_font
                    c.border = thin_border

            # Auto Column Widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row >= header_row:
                        val_str = str(cell.value or "")
                        max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            # Freeze Header Pane below row 4
            ws.freeze_panes = "A5"
            # Add AutoFilter
            ws.auto_filter.ref = f"A4:J{header_row + len(persons)}"

            wb.save(file_path)
            return True
        except Exception as e:
            log_error(f"Error exporting to Excel: {e}", exc_info=True)
            return False
